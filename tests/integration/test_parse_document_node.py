"""真实 MySQL 文件记录与实际 DOCX 联合验证解析节点及图中的状态传递。"""

import hashlib
from importlib import import_module
from uuid import uuid4

import pytest
from docx import Document

from app.graph.review_graph import ReviewNodes, build_review_graph
from app.models import AppUser, ContractFile, ContractTask
from app.parsing.exceptions import ParseBlocked

node_module = import_module("app.graph.nodes.parse_document")
pytestmark = pytest.mark.integration


@pytest.fixture
def node_context(factory, tmp_path, monkeypatch):
    path = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph(
        "甲方向乙方采购设备，双方按合同约定交付、付款和验收，并履行保密及违约责任。" * 3
    )
    document.save(path)
    with factory() as session:
        user = AppUser(username=f"node-{uuid4().hex}", display_name="节点测试", role_code="BIZ")
        session.add(user)
        session.flush()
        task = ContractTask(
            task_no=f"NODE-{uuid4().hex[:20]}", contract_name="节点测试合同", applicant_id=user.id
        )
        session.add(task)
        session.flush()
        file = ContractFile(
            task_id=task.id,
            original_name="contract.docx",
            storage_path=str(path),
            file_size=path.stat().st_size,
            file_md5=hashlib.md5(path.read_bytes(), usedforsecurity=False).hexdigest(),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_ext="docx",
        )
        session.add(file)
        session.commit()
        state = {"task_id": task.id, "file_id": file.id, "review_version": task.review_version}
    monkeypatch.setattr(node_module, "get_session_factory", lambda: factory)
    return state, path


def test_node_reads_record_and_returns_valid_document(node_context):
    state, _ = node_context
    update = node_module.parse_document(state)
    assert set(update) == {"document"}
    assert update["document"].file_id == state["file_id"]
    assert "采购设备" in update["document"].full_text
    assert "document" not in state  # 节点返回增量，不直接修改调用者的输入。


def test_graph_passes_parsed_document_to_next_node(node_context):
    state, _ = node_context

    def structure(current):
        assert current["document"].file_id == state["file_id"]
        assert current["document"].paragraphs
        return {"review_summary": "测试节点收到解析结果"}

    graph = build_review_graph(
        ReviewNodes(
            structure_document=structure,
            rule_review=lambda current: {},
            ai_review=lambda current: {},
            save_results=lambda current: {},
        )
    )
    result = graph.invoke(state)
    assert result["review_summary"] == "测试节点收到解析结果"
    assert result["document"].file_id == state["file_id"]


def test_file_cannot_be_read_under_another_task(node_context, factory):
    state, _ = node_context
    with factory() as session:
        owner = session.get(ContractTask, state["task_id"]).applicant_id
        other = ContractTask(
            task_no=f"OTHER-{uuid4().hex[:20]}", contract_name="其他任务", applicant_id=owner
        )
        session.add(other)
        session.commit()
        wrong_state = {**state, "task_id": other.id}
    with pytest.raises(ParseBlocked) as error:
        node_module.parse_document(wrong_state)
    assert error.value.reason == "ATTACHMENT_MISSING"


def test_deleted_task_is_rejected(node_context, factory):
    state, _ = node_context
    with factory() as session:
        session.get(ContractTask, state["task_id"]).is_deleted = True
        session.commit()
    with pytest.raises(ParseBlocked) as error:
        node_module.parse_document(state)
    assert error.value.reason == "ATTACHMENT_MISSING"


def test_outdated_review_version_is_rejected(node_context):
    state, _ = node_context
    with pytest.raises(ValueError, match="审查版本"):
        node_module.parse_document({**state, "review_version": state["review_version"] + 1})


def test_unsupported_file_format_is_rejected(node_context, factory):
    state, _ = node_context
    with factory() as session:
        session.get(ContractFile, state["file_id"]).file_ext = "pdf"
        session.commit()
    with pytest.raises(ParseBlocked) as error:
        node_module.parse_document(state)
    assert error.value.reason == "UNSUPPORTED_FORMAT"


def test_missing_original_file_preserves_failure_reason(node_context):
    state, path = node_context
    path.unlink()
    with pytest.raises(ParseBlocked) as error:
        node_module.parse_document(state)
    assert error.value.reason == "ATTACHMENT_MISSING"


def test_parser_failure_preserves_empty_body_reason(node_context):
    state, path = node_context
    Document().save(path)
    with pytest.raises(ParseBlocked) as error:
        node_module.parse_document(state)
    assert error.value.reason == "EMPTY_BODY"
