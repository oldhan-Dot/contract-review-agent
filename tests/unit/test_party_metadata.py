"""验证甲乙方主体提取、代码校验与原文证据，使用合成公司和测试代码。"""

import json

import pytest
from docx import Document

from app.parsing.docx_channel import parse_docx
from app.parsing.ir import ParsedDocument
from app.structuring.metadata import extract_party_metadata, validate_uscc

# 测试用合成代码：校验位通过不代表对应真实登记企业。
CODE_A = "91350211M000100Y46"
CODE_B = "913100001234567896"
NAME_A = "示例采购有限公司"
NAME_B = "示例供应有限公司"


def make_document(texts, *, table=False):
    paragraphs = []
    cursor = 0
    for index, text in enumerate(texts):
        paragraphs.append(
            {
                "pid": f"file-12:p{index}",
                "index": index,
                "text": text,
                "char_start": cursor,
                "char_end": cursor + len(text),
                "in_table": table,
                "table_id": "table-0" if table else None,
            }
        )
        cursor += len(text) + 1
    full_text = "\n".join(texts)
    return ParsedDocument(
        doc_id="file-12",
        file_id=12,
        source_format="DOCX",
        channel="docx",
        paragraphs=paragraphs,
        full_text=full_text,
        quality={
            "precision_cap": "EXACT",
            "char_count": len(full_text),
            "paragraph_count": len(texts),
        },
    )


def test_complete_party_sections_have_distinct_codes_and_valid_anchors():
    document = make_document(
        [
            f"甲方：{NAME_A}",
            f"统一社会信用代码：{CODE_A}",
            f"乙方：{NAME_B}",
            f"统一社会信用代码：{CODE_B}",
        ]
    )
    result = extract_party_metadata(document)
    expected = {
        "party_a_name": NAME_A,
        "party_a_uscc": CODE_A,
        "party_b_name": NAME_B,
        "party_b_uscc": CODE_B,
    }
    for key, value in expected.items():
        field = result.get(key)
        assert field.field_value == value
        assert not field.is_missing and field.is_required
        assert field.source == "auto" and field.confidence == 0.9
        assert field.anchor.matched_text == value
        field.anchor.validate_against_document(document)
    assert result.get("amount").is_missing is True
    assert len(result.items) == 15


def test_inline_party_sections_do_not_mix_names_or_codes():
    document = make_document(
        [f"甲方：{NAME_A} 统一社会信用代码：{CODE_A}；乙方：{NAME_B} 信用代码：{CODE_B}"]
    )
    result = extract_party_metadata(document)
    assert result.get("party_a_name").field_value == NAME_A
    assert result.get("party_b_name").field_value == NAME_B
    assert result.get("party_a_uscc").field_value == CODE_A
    assert result.get("party_b_uscc").field_value == CODE_B


def test_missing_a_code_does_not_reuse_b_code():
    document = make_document([f"甲方：{NAME_A}", f"乙方：{NAME_B}", f"信用代码：{CODE_B}"])
    result = extract_party_metadata(document)
    assert result.get("party_a_uscc").is_missing is True
    assert result.get("party_a_uscc").anchor is None
    assert result.get("party_b_uscc").field_value == CODE_B


def test_explicit_code_roles_override_nearest_party_section():
    document = make_document(
        [
            f"甲方：{NAME_A}",
            f"乙方：{NAME_B}",
            f"甲方统一社会信用代码：{CODE_A}",
            f"信用代码（乙方）：{CODE_B}",
        ]
    )
    result = extract_party_metadata(document)
    assert result.get("party_a_uscc").field_value == CODE_A
    assert result.get("party_b_uscc").field_value == CODE_B
    assert result.get("party_b_name").field_value == NAME_B


def test_code_role_annotation_does_not_become_company_name():
    result = extract_party_metadata(make_document([f"信用代码（甲方）：{CODE_A}"]))
    assert result.get("party_a_name").is_missing is True
    assert result.get("party_a_uscc").field_value == CODE_A


def test_code_without_role_context_is_not_guessed():
    result = extract_party_metadata(make_document([f"统一社会信用代码：{CODE_A}"]))
    assert result.get("party_a_uscc").is_missing is True
    assert result.get("party_b_uscc").is_missing is True


def test_numbered_body_stops_party_code_context():
    document = make_document([f"甲方：{NAME_A}", "第一条 开票要求", f"信用代码：{CODE_B}"])
    assert extract_party_metadata(document).get("party_a_uscc").is_missing is True


@pytest.mark.parametrize("label", ["甲方（买方）", "买方（甲方）", "甲 方", "甲方名称"])
def test_explicit_party_label_variants(label):
    result = extract_party_metadata(make_document([f"{label}：{NAME_A}"]))
    assert result.get("party_a_name").field_value == NAME_A


def test_bare_role_and_separate_name_label_are_supported():
    document = make_document(["甲方", "单位名称：", NAME_A, "统一社会信用代码", CODE_A])
    result = extract_party_metadata(document)
    assert result.get("party_a_name").field_value == NAME_A
    assert result.get("party_a_uscc").field_value == CODE_A


def test_buyer_seller_labels_without_explicit_roles_are_not_assumed():
    result = extract_party_metadata(make_document([f"买方：{NAME_A}", f"卖方：{NAME_B}"]))
    assert result.get("party_a_name").is_missing is True
    assert result.get("party_b_name").is_missing is True


def test_invalid_name_candidate_is_retained_with_warning():
    field = extract_party_metadata(make_document(["甲方：张三"])).get("party_a_name")
    assert field.field_value == "张三" and field.field_value_raw == "张三"
    assert field.confidence == 0.3 and field.warnings


@pytest.mark.parametrize("placeholder", ["______", "待填写", "（盖章）"])
def test_placeholder_names_remain_missing(placeholder):
    result = extract_party_metadata(make_document([f"甲方：{placeholder}", f"乙方：{NAME_B}"]))
    assert result.get("party_a_name").is_missing is True
    assert result.get("party_b_name").field_value == NAME_B


@pytest.mark.parametrize(
    "code", ["123", "91350211M000100Y43", "91350211I000100Y46", "9135-0211M000100Y46"]
)
def test_invalid_codes_are_kept_with_low_confidence(code):
    field = extract_party_metadata(make_document([f"甲方：{NAME_A}", f"信用代码：{code}"])).get(
        "party_a_uscc"
    )
    assert field.field_value == code.upper() and field.field_value_raw == code
    assert field.confidence == 0.3 and field.warnings
    assert not field.is_missing


def test_lowercase_code_is_normalized_but_raw_evidence_is_preserved():
    document = make_document([f"甲方：{NAME_A}", f"信用代码：{CODE_A.lower()}"])
    field = extract_party_metadata(document).get("party_a_uscc")
    assert field.field_value == CODE_A
    assert field.field_value_raw == CODE_A.lower()
    assert field.anchor.matched_text == CODE_A.lower()
    assert field.confidence == 0.9


def test_multiple_candidates_use_first_position_and_keep_alternatives():
    document = make_document([f"甲方：{NAME_A}", f"甲方：{NAME_B}"])
    field = extract_party_metadata(document).get("party_a_name")
    assert field.field_value == NAME_A
    assert json.loads(field.field_value_raw) == [NAME_A, NAME_B]
    assert field.warnings and field.anchor.matched_text == NAME_A


def test_identical_repeated_candidates_are_not_reported_as_conflict():
    field = extract_party_metadata(make_document([f"甲方：{NAME_A}", f"甲方：{NAME_A}"])).get(
        "party_a_name"
    )
    assert field.field_value_raw == NAME_A
    assert field.warnings == []


def test_oversized_name_does_not_abort_extraction():
    name = "某" * 1100 + "有限公司"
    result = extract_party_metadata(make_document([f"甲方：{name}", f"乙方：{NAME_B}"]))
    field = result.get("party_a_name")
    assert field.is_missing is True and field.anchor is None
    assert len(field.field_value_raw) == 1024
    assert field.warnings
    assert result.get("party_b_name").field_value == NAME_B


def test_many_candidates_keep_bounded_raw_list_without_aborting():
    names = [f"示例第{index}家" + "科技" * 20 + "有限公司" for index in range(30)]
    field = extract_party_metadata(make_document([f"甲方：{name}" for name in names])).get(
        "party_a_name"
    )
    assert field.field_value == names[0]
    assert len(field.field_value_raw) <= 1024
    assert 0 < len(json.loads(field.field_value_raw)) < len(names)
    assert any("未全部展开" in warning for warning in field.warnings)


def test_blank_credit_code_placeholder_stays_missing():
    result = extract_party_metadata(make_document([f"甲方：{NAME_A}", "信用代码：______"]))
    assert result.get("party_a_uscc").is_missing is True


def test_empty_document_returns_missing_fields_without_blocking():
    result = extract_party_metadata(make_document([]))
    assert all(
        result.get(key).is_missing
        for key in ("party_a_name", "party_b_name", "party_a_uscc", "party_b_uscc")
    )


def test_parallel_table_headers_do_not_assign_a_values_to_b():
    document = make_document(
        ["项目", "甲方", "乙方", "名称", NAME_A, NAME_B, "信用代码", CODE_A, CODE_B], table=True
    )
    result = extract_party_metadata(document)
    for key in ("party_a_name", "party_b_name", "party_a_uscc", "party_b_uscc"):
        assert result.get(key).is_missing is True
        assert result.get(key).warnings


def test_actual_docx_vertical_table_is_extracted_with_real_positions(tmp_path):
    word = Document()
    word.add_paragraph("采购合同的主体信息，以本表双方名称和代码为准。")
    table = word.add_table(rows=4, cols=2)
    for row, values in zip(
        table.rows,
        [
            ("甲方", NAME_A),
            ("统一社会信用代码", CODE_A),
            ("乙方", NAME_B),
            ("统一社会信用代码", CODE_B),
        ],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / "contract.docx"
    word.save(path)
    document = parse_docx(path, file_id=12)
    result = extract_party_metadata(document)
    assert result.get("party_a_name").field_value == NAME_A
    assert result.get("party_b_uscc").field_value == CODE_B
    assert result.get("party_b_uscc").anchor.paragraph_index > 0
    result.validate_against_document(document)


@pytest.mark.parametrize("code", [CODE_A, CODE_B, CODE_A.lower(), f" {CODE_B} "])
def test_known_checksum_results_are_accepted(code):
    assert validate_uscc(code) is True


@pytest.mark.parametrize("code", ["", "123", CODE_A[:-1] + "3", "I" * 18])
def test_invalid_checksum_or_character_set_is_rejected(code):
    assert validate_uscc(code) is False
