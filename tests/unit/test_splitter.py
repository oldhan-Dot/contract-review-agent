"""用独立合同片段核对切分粒度、编号层级和原文定位。"""

import pytest
from docx import Document

from app.parsing.docx_channel import parse_docx
from app.parsing.exceptions import ParseBlocked
from app.parsing.ir import ParsedDocument
from app.structuring.splitter import match_clause_no, split_clauses


def make_document(texts, *, headings=None, bold_indexes=(), table_indexes=(), page_numbers=None):
    paragraphs = []
    cursor = 0
    for index, text in enumerate(texts):
        paragraphs.append(
            {
                "pid": f"file-1:p{index}",
                "index": index,
                "text": text,
                "char_start": cursor,
                "char_end": cursor + len(text),
                "heading_level": (headings or {}).get(index),
                "runs": [{"text": text, "bold": True}] if index in bold_indexes else [],
                "in_table": index in table_indexes,
                "table_id": "table-0" if index in table_indexes else None,
                "page_no": page_numbers[index] if page_numbers else None,
            }
        )
        cursor += len(text) + 1
    return ParsedDocument(
        doc_id="file-1",
        file_id=1,
        source_format="PDF_TEXT" if page_numbers else "DOCX",
        channel="pdf_text" if page_numbers else "docx",
        paragraphs=paragraphs,
        full_text="\n".join(texts),
        quality={"precision_cap": "EXACT", "paragraph_count": len(texts)},
    )


def test_two_numbered_clauses_keep_preamble_and_table_text():
    document = make_document(
        [
            "采购合同",
            " 第一条 付款  ",
            "甲方支付货款。",
            "设备名称",
            "设备A",
            "第二条 交付",
            "乙方负责交付。",
        ],
        table_indexes=(3, 4),
    )
    clauses = split_clauses(document)
    assert [c.clause_no for c in clauses] == [None, "第一条", "第二条"]
    assert clauses[1].text == "付款  \n甲方支付货款。\n设备名称\n设备A"
    assert clauses[1].source_text(document).startswith(" 第一条")
    assert clauses[1].paragraph_index_end == 4
    assert all(c.split_strategy == "NUMBER" for c in clauses)
    assert [c.index for c in clauses] == [1, 2, 3]


@pytest.mark.parametrize(
    "prefix,level",
    [
        ("第一条", 1),
        ("第 12 条", 1),
        ("第三条之一", 1),
        ("一、", 1),
        ("1.1", 1),
        ("1.2", 1),
        ("1.1.1", 2),
        ("1）", 1),
        ("（一）", 2),
        ("(1)", 3),
        ("【第三条】", 1),
    ],
)
def test_supported_number_prefixes(prefix, level):
    result = match_clause_no(prefix + " 付款约定")
    assert result is not None
    assert result[:2] == (level, prefix)


@pytest.mark.parametrize("text", ["2026年合同", "1.25元", "1.25 万元", "甲方支付1.1万元"])
def test_amounts_and_dates_are_not_number_prefixes(text):
    assert match_clause_no(text) is None


def test_numbered_hierarchy_links_to_existing_parent():
    clauses = split_clauses(
        make_document(["第一条 付款", "（一）预付款", "(1) 支付比例", "第二条 交付"])
    )
    assert [c.level for c in clauses] == [1, 2, 3, 1]
    assert [c.parent_index for c in clauses] == [None, 1, 2, None]


def test_standalone_number_is_joined_with_following_body():
    clauses = split_clauses(make_document(["第一条", "甲方按时付款。", "第二条", "乙方按时交货。"]))
    assert [c.clause_no for c in clauses] == ["第一条", "第二条"]
    assert clauses[0].text == "甲方按时付款。"


def test_style_strategy_groups_body_until_next_heading():
    document = make_document(
        ["付款", "甲方支付货款。", "交付", "乙方交付设备。"], headings={0: 1, 2: 1}
    )
    clauses = split_clauses(document)
    assert [c.text for c in clauses] == ["付款\n甲方支付货款。", "交付\n乙方交付设备。"]
    assert all(c.split_strategy == "STYLE" for c in clauses)


def test_number_strategy_takes_priority_without_mixing_styles():
    document = make_document(
        ["第一条 付款", "补充标题", "付款义务。", "第二条 交付"], headings={1: 1}
    )
    clauses = split_clauses(document)
    assert len(clauses) == 2
    assert "补充标题" in clauses[0].text
    assert all(c.split_strategy == "NUMBER" for c in clauses)


def test_bold_pseudo_heading_groups_body():
    document = make_document(
        ["付款约定", "按约定支付。", "交付约定", "按约定交付。"], bold_indexes=(0, 2)
    )
    clauses = split_clauses(document)
    assert len(clauses) == 2
    assert all(c.split_strategy == "PSEUDO_HEADING" for c in clauses)


def test_natural_paragraph_fallback_preserves_short_obligations():
    clauses = split_clauses(make_document(["应付款", "不得转让。", "乙方应保密。 "]))
    assert [c.text for c in clauses] == ["应付款", "不得转让。", "乙方应保密。"]
    assert all(c.split_strategy == "BLANK_LINE" and c.level == 0 for c in clauses)


def test_long_clause_children_have_valid_parent_and_sentence_boundaries():
    document = make_document(["第一条 付款", "甲方应按时支付货款。" * 350, "第二条 交付"])
    clauses = split_clauses(document)
    children = [c for c in clauses if c.part_no is not None]
    assert len(children) > 1
    assert all(c.level == 0 and c.parent_index == 1 for c in children)
    assert [c.part_no for c in children] == list(range(1, len(children) + 1))
    assert all(len(c.source_text(document)) <= 800 for c in children)
    assert children[-1].char_end == clauses[0].char_end
    assert clauses[-1].clause_no == "第二条"
    for clause in clauses:
        clause.validate_against_document(document)


def test_long_unpunctuated_paragraph_is_hard_split_without_losing_characters():
    document = make_document(["甲" * 3201])
    clauses = split_clauses(document)
    parent, children = clauses[0], clauses[1:]
    assert len(children) == 5
    assert "".join(c.source_text(document) for c in children) == parent.source_text(document)
    assert all(len(c.text) <= 800 for c in children)


def test_page_boundary_does_not_split_one_clause():
    document = make_document(
        ["第一条 付款", "付款内容第一页。", "付款内容第二页。", "第二条 交付"],
        page_numbers=[1, 1, 2, 2],
    )
    clauses = split_clauses(document)
    assert len(clauses) == 2
    assert clauses[0].is_cross_page() is True
    assert (clauses[0].page_start, clauses[0].page_end) == (1, 2)


def test_empty_document_is_blocked():
    with pytest.raises(ParseBlocked) as error:
        split_clauses(make_document([]))
    assert error.value.reason == "EMPTY_BODY"


def test_invalid_length_limits_are_rejected():
    with pytest.raises(ValueError):
        split_clauses(make_document(["付款条款"]), max_len=300, target_len=800)


def test_real_docx_parser_output_can_be_split(tmp_path):
    word = Document()
    word.add_paragraph("采购合同")
    word.add_heading("第一条 付款", level=1)
    word.add_paragraph("甲方应在验收合格后按约定支付合同价款，并由乙方提供合法有效的发票。" * 2)
    word.add_table(rows=1, cols=1).cell(0, 0).text = "付款账户信息"
    word.add_heading("第二条 交付", level=1)
    word.add_paragraph("乙方应在约定期限内交付合格设备。")
    path = tmp_path / "contract.docx"
    word.save(path)
    document = parse_docx(path, file_id=1)
    clauses = split_clauses(document)
    assert [c.clause_no for c in clauses] == [None, "第一条", "第二条"]
    assert "付款账户信息" in clauses[1].text
    for clause in clauses:
        clause.validate_against_document(document)
