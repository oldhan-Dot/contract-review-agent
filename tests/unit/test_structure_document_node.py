"""结构化节点的真实处理、状态传递和异常中止，不依赖数据库或模型。"""

import json
from importlib import import_module

import pytest
from docx import Document

from app.graph.review_graph import ReviewNodes, build_review_graph
from app.parsing.docx_channel import parse_docx
from app.parsing.exceptions import ParseBlocked
from app.parsing.ir import ParsedDocument, SourceFormat
from app.structuring.metadata_models import MetadataResult
from app.structuring.models import ClauseType

node = import_module("app.graph.nodes.structure_document")


def document(*texts):
    paragraphs, cursor = [], 0
    for index, text in enumerate(texts):
        paragraphs.append(
            dict(
                pid=f"file-12:p{index}",
                index=index,
                text=text,
                char_start=cursor,
                char_end=cursor + len(text),
            )
        )
        cursor += len(text) + 1
    full_text = "\n".join(texts)
    return ParsedDocument(
        doc_id="file-12",
        file_id=12,
        source_format="DOCX",
        channel="docx",
        paragraphs=paragraphs,
        full_text=full_text,
        quality=dict(precision_cap="EXACT", char_count=len(full_text), paragraph_count=len(texts)),
    )


def state(parsed=None):
    return dict(
        task_id=1,
        file_id=12,
        review_version=1,
        document=parsed or document("设备采购合同", "第一条 保密义务", "不得披露商业秘密。"),
    )


def test_real_split_classify_extract_sequence_and_input_unchanged(monkeypatch):
    current = state()
    before = current["document"].model_dump()
    calls = []
    for name in ("split_clauses", "classify_clauses", "extract_metadata"):
        original = getattr(node, name)

        def track(*args, _name=name, _original=original, **kwargs):
            calls.append(_name)
            return _original(*args, **kwargs)

        monkeypatch.setattr(node, name, track)
    result = node.structure_document(current)
    assert calls == ["split_clauses", "classify_clauses", "extract_metadata"]
    assert set(result) == {"clauses", "metadata"}
    assert any(item.clause_type == ClauseType.CONFIDENTIALITY for item in result["clauses"])
    assert result["metadata"].get("contract_title").field_value == "设备采购合同"
    assert isinstance(result["metadata"], MetadataResult)
    for clause in result["clauses"]:
        clause.validate_against_document(current["document"])
    result["metadata"].validate_against_document(current["document"])
    assert current["document"].model_dump() == before
    assert "clauses" not in current and "metadata" not in current


@pytest.mark.parametrize("value", [None, {}, "文档正文"])
def test_missing_or_wrong_document_rejected_before_processing(monkeypatch, value):
    def forbidden(*args, **kwargs):
        raise AssertionError("非法输入不能进入切分")

    monkeypatch.setattr(node, "split_clauses", forbidden)
    with pytest.raises(ValueError, match="document"):
        node.structure_document({**state(), "document": value})


@pytest.mark.parametrize("value", [True, 0, "12", 2**64, 13])
def test_invalid_or_different_file_id_rejected(value):
    with pytest.raises(ValueError, match="file_id"):
        node.structure_document({**state(), "file_id": value})


def test_unsupported_format_explicitly_blocked():
    parsed = document("测试 PDF 合同").model_copy(update={"source_format": SourceFormat.PDF_TEXT})
    with pytest.raises(ParseBlocked) as error:
        node.structure_document(state(parsed))
    assert error.value.reason == "UNSUPPORTED_FORMAT"


def test_invalid_task_name_rejected():
    with pytest.raises(ValueError, match="task_name"):
        node.structure_document({**state(), "task_name": 123})


def test_explicit_task_name_fallback_forwarded():
    result = node.structure_document(
        {**state(document("甲方：示例采购有限公司")), "task_name": "采购审批任务"}
    )
    title = result["metadata"].get("contract_title")
    assert title.field_value == "采购审批任务"
    assert title.source == "auto" and title.anchor is None


def test_missing_fields_and_unresolved_classification_do_not_block():
    result = node.structure_document(state(document("双方另行商议具体安排。")))
    assert result["metadata"].get("amount").is_missing
    assert result["clauses"][0].clause_type == ClauseType.OTHER
    assert json.loads(result["clauses"][0].classify_rule_hit)["needs_fallback"] is True


def test_long_clause_preserves_parent_children_and_source_ranges():
    parsed = document("第一条 保密义务", "不得披露商业秘密。" * 400)
    result = node.structure_document(state(parsed))
    children = [item for item in result["clauses"] if item.part_no is not None]
    assert children
    by_index = {item.index: item for item in result["clauses"]}
    for child in children:
        assert child.parent_index in by_index
        assert child.clause_type == ClauseType.CONFIDENTIALITY
        child.validate_against_document(parsed)


def test_failure_does_not_run_extraction_or_write_partial_state(monkeypatch):
    current = state()

    def fail(*args, **kwargs):
        raise ValueError("模拟归类异常")

    def forbidden(*args, **kwargs):
        raise AssertionError("归类失败后不能继续字段提取")

    monkeypatch.setattr(node, "classify_clauses", fail)
    monkeypatch.setattr(node, "extract_metadata", forbidden)
    with pytest.raises(ValueError, match="模拟归类异常"):
        node.structure_document(current)
    assert "clauses" not in current and "metadata" not in current


def test_empty_document_preserves_parse_blocked_reason():
    parsed = document("   ")
    with pytest.raises(ParseBlocked) as error:
        node.structure_document(state(parsed))
    assert error.value.reason == "EMPTY_BODY"


def test_graph_passes_typed_results_and_rerun_replaces_lists():
    parsed = document("设备采购合同", "第一条 保密义务", "不得披露商业秘密。")
    calls = []

    def downstream(current):
        assert isinstance(current["document"], ParsedDocument)
        assert isinstance(current["metadata"], MetadataResult)
        assert current["task_name"] == "审批任务"
        assert current["review_degraded"] is True
        assert current["review_degraded_reason"] == "已有降级原因"
        assert any(item.clause_type == ClauseType.CONFIDENTIALITY for item in current["clauses"])
        calls.append("rules")
        return {}

    graph = build_review_graph(
        ReviewNodes(
            parse_document=lambda current: {"document": parsed},
            rule_review=downstream,
            ai_review=lambda current: {},
            save_results=lambda current: {},
        )
    )
    initial = dict(
        task_id=1,
        file_id=12,
        review_version=1,
        task_name="审批任务",
        review_degraded=True,
        review_degraded_reason="已有降级原因",
    )
    first = graph.invoke(initial)
    second = graph.invoke(first)
    assert calls == ["rules", "rules"]
    assert len(second["clauses"]) == len(first["clauses"])
    assert second["metadata"].model_dump() == first["metadata"].model_dump()


def test_real_docx_parser_to_structure_node(tmp_path):
    word = Document()
    for text in [
        "设备采购合同",
        "合同编号：HT-2026/001",
        "甲方：示例采购有限公司",
        "第一条 金额与付款",
        "合同金额：人民币1000元，按照约定支付货款。",
        "第二条 保密义务",
        "不得披露商业秘密。",
    ]:
        word.add_paragraph(text)
    path = tmp_path / "structure.docx"
    word.save(path)
    parsed = parse_docx(path, file_id=12)
    result = node.structure_document(state(parsed))
    assert result["metadata"].get("contract_no").field_value == "HT-2026/001"
    assert result["metadata"].get("party_a_name").field_value == "示例采购有限公司"
    assert result["metadata"].get("amount").field_value == "1000.00"
    assert any(item.clause_type == ClauseType.PAYMENT for item in result["clauses"])
    assert any(item.clause_type == ClauseType.CONFIDENTIALITY for item in result["clauses"])
    result["metadata"].validate_against_document(parsed)
