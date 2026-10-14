"""使用实际生成的 Word 文件验证顺序、样式、表格和失败原因。"""

import pytest
from docx import Document
from docx.shared import Pt
from msoffcrypto.format.ooxml import OOXMLFile

from app.parsing.docx_channel import parse_docx
from app.parsing.exceptions import ParseBlocked

BODY = "甲方向乙方采购设备，乙方应当按合同约定交付，双方应当按约定履行付款和验收义务。"


def test_body_table_order_styles_and_offsets(tmp_path):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "页眉公司名称"
    document.sections[0].footer.paragraphs[0].text = "页脚公司名称"
    document.add_heading("采购合同", level=1)
    before = f"  {BODY}  "
    paragraph = document.add_paragraph(before)
    run = paragraph.add_run("重点约定")
    run.bold = True
    run.italic = True
    run.font.name = "宋体"
    run.font.size = Pt(14)
    document.add_paragraph("   ")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "设备名称"
    table.cell(0, 1).text = "测试设备"
    document.add_paragraph("表格后的付款条款")
    path = tmp_path / "contract.docx"
    document.save(path)

    result = parse_docx(path, file_id=12)
    assert [p.text for p in result.paragraphs] == [
        "采购合同",
        before + "重点约定",
        "设备名称",
        "测试设备",
        "表格后的付款条款",
    ]
    assert result.paragraphs[0].heading_level == 1
    assert result.paragraphs[0].style_name == "Heading 1"
    assert result.paragraphs[2].in_table is True
    assert result.paragraphs[2].table_id == result.paragraphs[3].table_id
    assert result.paragraphs[4].in_table is False
    styled_run = result.paragraphs[1].runs[-1]
    assert styled_run.bold is True and styled_run.italic is True
    assert styled_run.font_name == "宋体" and styled_run.font_size_pt == 14
    assert "页眉公司名称" not in result.full_text and "页脚公司名称" not in result.full_text
    assert all(p.page_no is None for p in result.paragraphs)
    assert result.quality.char_count == len(result.full_text)
    assert result.quality.paragraph_count == 5
    assert result.quality.empty_paragraph_ratio > 0
    for paragraph in result.paragraphs:
        assert result.slice_text(paragraph.char_start, paragraph.char_end) == paragraph.text


def test_merged_cells_are_read_once_and_nested_table_keeps_order(tmp_path):
    document = Document()
    document.add_paragraph(BODY * 2)
    table = document.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(1, 1))
    merged.text = "外层开始"
    nested = merged.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "内层条款"
    merged.add_paragraph("外层结束")
    path = tmp_path / "merged.docx"
    document.save(path)

    result = parse_docx(path, file_id=1)
    assert [p.text for p in result.paragraphs][1:] == ["外层开始", "内层条款", "外层结束"]
    assert result.paragraphs[1].table_id == result.paragraphs[3].table_id
    assert result.paragraphs[2].table_id != result.paragraphs[1].table_id


def test_heading_levels_match_prd_example(tmp_path):
    document = Document()
    for index in range(5):
        document.add_heading(f"一级标题 {index}", level=1)
    for index in range(18):
        document.add_heading(f"二级标题 {index}", level=2)
    path = tmp_path / "headings.docx"
    document.save(path)
    result = parse_docx(path, file_id=1)
    assert sum(p.heading_level == 1 for p in result.paragraphs) == 5
    assert sum(p.heading_level == 2 for p in result.paragraphs) == 18
    assert [p.index for p in result.paragraphs] == list(range(23))


@pytest.mark.parametrize("text", ["", "只有很短的文字"])
def test_empty_or_short_document_is_blocked(tmp_path, text):
    document = Document()
    document.add_paragraph(text)
    path = tmp_path / "empty.docx"
    document.save(path)
    with pytest.raises(ParseBlocked) as error:
        parse_docx(path, file_id=1)
    assert error.value.reason == "EMPTY_BODY"


def test_corrupt_file_is_blocked(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_bytes(b"This is not a Word document")
    with pytest.raises(ParseBlocked) as error:
        parse_docx(path, file_id=1)
    assert error.value.reason == "ATTACHMENT_CORRUPT"


def test_missing_file_is_blocked(tmp_path):
    with pytest.raises(ParseBlocked) as error:
        parse_docx(tmp_path / "missing.docx", file_id=1)
    assert error.value.reason == "ATTACHMENT_MISSING"


def test_real_encrypted_docx_is_blocked(tmp_path):
    document = Document()
    document.add_paragraph(BODY * 2)
    plain = tmp_path / "plain.docx"
    document.save(plain)
    encrypted = tmp_path / "encrypted.docx"
    with plain.open("rb") as source, encrypted.open("wb") as target:
        OOXMLFile(source).encrypt("test-password", target)
    with pytest.raises(ParseBlocked) as error:
        parse_docx(encrypted, file_id=1)
    assert error.value.reason == "ENCRYPTED_FILE"


def test_empty_heading_is_skipped_with_warning(tmp_path):
    document = Document()
    document.add_heading("", level=1)
    document.add_paragraph(BODY * 2)
    path = tmp_path / "empty-heading.docx"
    document.save(path)
    result = parse_docx(path, file_id=1)
    assert len(result.paragraphs) == 1
    assert result.paragraphs[0].index == 0
    assert result.quality.warnings


@pytest.mark.parametrize("file_id", [0, True, 2**64])
def test_invalid_file_id_is_rejected_before_reading(tmp_path, file_id):
    with pytest.raises(ValueError, match="file_id"):
        parse_docx(tmp_path / "not-opened.docx", file_id=file_id)
