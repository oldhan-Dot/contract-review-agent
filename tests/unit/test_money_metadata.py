"""金额提取的业务边界和实际 DOCX 原文定位。"""

import json

import pytest
from docx import Document

from app.parsing.docx_channel import parse_docx
from app.parsing.ir import ParsedDocument
from app.structuring.metadata import extract_metadata
from app.structuring.money_metadata import extract_money_metadata


def document(text):
    return ParsedDocument(
        doc_id="file-12",
        file_id=12,
        source_format="DOCX",
        channel="docx",
        paragraphs=[dict(pid="file-12:p0", index=0, text=text, char_start=0, char_end=len(text))],
        full_text=text,
        quality=dict(precision_cap="EXACT", char_count=len(text), paragraph_count=1),
    )


def extract(text):
    return extract_money_metadata(document(text))


def test_prd_example():
    text = "合同总金额：人民币壹佰贰拾万元整（¥1,200,000.00）"
    result = extract(text)
    assert result.get("amount").field_value == "1200000.00"
    assert result.get("amount_in_words_value").field_value == "1200000.00"
    assert result.get("currency").field_value == "CNY"
    assert result.get("amount").anchor.matched_text == "1,200,000.00"
    assert result.get("amount").warnings == []


@pytest.mark.parametrize(
    "text,value,code",
    [
        ("合同金额：USD 100.50", "100.50", "USD"),
        ("合同金额：€100", "100.00", "EUR"),
        ("合同金额：HK$100", "100.00", "HKD"),
        ("合同金额：100港元", "100.00", "HKD"),
        ("合同金额：￥1.2万元", "12000.00", "CNY"),
        ("合同总额：壹万元整", "10000.00", "CNY"),
        ("合同金额：１００．５０元", "100.50", "CNY"),
        ("合同金额：100", "100.00", None),
        ("合同金额：100日元", "100.00", None),
    ],
)
def test_amount_and_local_currency(text, value, code):
    result = extract(text)
    assert result.get("amount").field_value == value
    assert result.get("currency").field_value == code


@pytest.mark.parametrize(
    "text",
    [
        "日期：2026年10月15日，比例30%，税号913100001234567896",
        "预付款：￥200；单价：￥10",
        "本合同采用一元化管理",
    ],
)
def test_unrelated_numbers_are_not_contract_total(text):
    assert extract(text).get("amount").is_missing


def test_total_beats_earlier_component_and_currency():
    result = extract("预付款：USD200；合同总额：人民币1000元")
    assert result.get("amount").field_value == "1000.00"
    assert result.get("currency").field_value == "CNY"


def test_conflict_retains_both_and_numeric_wins():
    result = extract("合同金额：￥100（大写：贰佰元整）")
    assert result.get("amount").field_value == "100.00"
    assert result.get("amount_in_words_value").field_value == "200.00"
    assert any("不一致" in warning for warning in result.get("amount").warnings)


def test_multiple_totals_are_not_silently_overwritten():
    field = extract("合同金额：￥100；合同金额：￥200").get("amount")
    assert field.field_value == "100.00"
    assert field.confidence == 0.3
    assert len(json.loads(field.field_value_raw)) == 2


@pytest.mark.parametrize("raw", ["0", "-10", "1000000000001", "1,20", "10.123"])
def test_invalid_observation_is_preserved(raw):
    field = extract(f"合同金额：￥{raw}").get("amount")
    assert not field.is_missing
    assert field.warnings
    assert field.confidence == 0.3
    assert field.anchor.matched_text == raw


def test_large_candidate_list_still_has_valid_bounded_json():
    field = extract("；".join(f"合同金额：￥{i}" for i in range(1, 100))).get("amount")
    assert len(field.field_value_raw) <= 1024
    assert json.loads(field.field_value_raw)


def test_real_docx_merge_preserves_party_and_evidence(tmp_path):
    word = Document()
    word.add_paragraph(
        "采购合同：双方根据约定履行交付、验收及付款义务，本合同金额以以下总额约定为准。"
    )
    word.add_paragraph("甲方：示例采购有限公司")
    word.add_paragraph("合同总额：￥1,200.00")
    path = tmp_path / "money.docx"
    word.save(path)
    parsed = parse_docx(path, file_id=12)
    result = extract_metadata(parsed)
    assert result.get("party_a_name").field_value == "示例采购有限公司"
    assert result.get("amount").field_value == "1200.00"
    assert result.get("sign_date").is_missing
    result.validate_against_document(parsed)


def test_explicit_currency_declaration():
    result = extract("币种：USD；合同金额：100")
    assert result.get("currency").field_value == "USD"
    assert result.get("currency").anchor.matched_text == "USD"


@pytest.mark.parametrize(
    "text", ["合同金额：见第101条", "合同金额：附件编号123", "合同金额：待填写，2026年签署"]
)
def test_references_are_not_amounts(text):
    assert extract(text).get("amount").is_missing


def test_unspaced_currency_code():
    result = extract("合同总额：USD100")
    assert result.get("amount").field_value == "100.00"
    assert result.get("currency").field_value == "USD"


def test_unrelated_chinese_amount_is_not_compared():
    result = extract("合同金额：￥100；预付款：贰佰元整")
    assert result.get("amount").warnings == []
    assert result.get("amount_in_words").is_missing
