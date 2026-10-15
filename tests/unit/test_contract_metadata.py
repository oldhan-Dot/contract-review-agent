"""验证名称/编号提取、兜底来源、原文位置和组合提取。"""

import json
from datetime import date

import pytest
from docx import Document

from app.parsing.docx_channel import parse_docx
from app.parsing.ir import ParsedDocument
from app.structuring.contract_metadata import extract_contract_metadata
from app.structuring.metadata import extract_metadata


def document(*texts):
    paragraphs, cursor = [], 0
    for index, text in enumerate(texts):
        paragraphs.append(
            dict(
                pid=f"file-12:p{index}",
                index=index,
                text=text,
                char_start=cursor,
                char_end=cursor + len(text),
            )
        )
        cursor += len(text) + 1
    full_text = "\n".join(texts)
    return ParsedDocument(
        doc_id="file-12",
        file_id=12,
        source_format="DOCX",
        channel="docx",
        paragraphs=paragraphs,
        full_text=full_text,
        quality=dict(precision_cap="EXACT", char_count=len(full_text), paragraph_count=len(texts)),
    )


@pytest.mark.parametrize(
    "title",
    ["设备采购合同", "技术服务协议", "设备买卖契约", "设备采购合同（副本）", "设备采购合同书"],
)
def test_plain_cover_title(title):
    field = extract_contract_metadata(document(f"  {title}  ")).get("contract_title")
    assert field.field_value == title
    assert field.anchor.matched_text == title
    assert field.anchor.char_start == 2
    assert field.confidence == 0.9


def test_explicit_title_and_number_on_same_line():
    parsed = document("合同名称：设备采购合同 合同编号：HT-2026/001")
    result = extract_contract_metadata(parsed)
    assert result.get("contract_title").field_value == "设备采购合同"
    assert result.get("contract_no").field_value == "HT-2026/001"
    result.validate_against_document(parsed)


def test_label_value_in_next_paragraph():
    result = extract_contract_metadata(
        document("合同名称：", "设备采购合同", "合同编号：", "HT-2026_001")
    )
    assert result.get("contract_title").anchor.paragraph_index == 1
    assert result.get("contract_no").anchor.paragraph_index == 3


@pytest.mark.parametrize("label", ["合同编号", "合同编码", "编号"])
def test_supported_number_labels(label):
    field = extract_contract_metadata(document(f"{label}：Ab-2026/001_x")).get("contract_no")
    assert field.field_value == "Ab-2026/001_x"
    assert field.confidence == 0.9


@pytest.mark.parametrize(
    "text",
    [
        "发票编号：FP-2026",
        "订单编号：DD-2026",
        "银行账号：1234567890",
        "原合同编号：OLD-2026",
        "主合同编号：MAIN-2026",
        "关联合同编号：REF-2026",
        "甲方：示例公司",
    ],
)
def test_other_identifiers_not_contract_number(text):
    assert extract_contract_metadata(document(text)).get("contract_no").is_missing


@pytest.mark.parametrize("raw", ["AB", "合同2026", "A" * 33, "ＡＢＣ１２３", "AB@2026"])
def test_invalid_explicit_number_preserved_without_partial_match(raw):
    field = extract_contract_metadata(document(f"合同编号：{raw}")).get("contract_no")
    assert field.field_value == raw
    assert field.anchor.matched_text == raw
    assert field.confidence == 0.3
    assert field.warnings


@pytest.mark.parametrize("raw", ["____", "待填写", "暂无", "---"])
def test_placeholder_number_is_missing(raw):
    assert extract_contract_metadata(document(f"合同编号：{raw}")).get("contract_no").is_missing


def test_blank_number_does_not_swallow_next_label():
    field = extract_contract_metadata(document("合同编号：", "甲方：示例采购有限公司")).get(
        "contract_no"
    )
    assert field.is_missing


def test_multiple_different_numbers_preserve_candidates():
    field = extract_contract_metadata(document("合同编号：HT-0001", "合同编号：HT-0002")).get(
        "contract_no"
    )
    assert field.field_value == "HT-0001"
    assert field.confidence == 0.3
    assert len(json.loads(field.field_value_raw)) == 2


def test_repeated_same_number_does_not_warn():
    field = extract_contract_metadata(document("合同编号：HT-0001", "合同编号：HT-0001")).get(
        "contract_no"
    )
    assert field.confidence == 0.9
    assert field.warnings == []


def test_title_search_stops_at_cover_boundary():
    result = extract_contract_metadata(document(*(["正文内容"] * 20), "设备采购合同"))
    assert result.get("contract_title").is_missing


def test_title_search_stops_at_character_boundary():
    assert (
        extract_contract_metadata(document("前言" * 1000, "设备采购合同"))
        .get("contract_title")
        .is_missing
    )


@pytest.mark.parametrize(
    "text",
    [
        "根据相关法律双方订立本合同。",
        "第一条 本合同",
        "合同编号：HT-0001",
        "第一条本合同",
        "甲乙双方同意订立本合同",
    ],
)
def test_body_text_is_not_a_cover_title(text):
    assert extract_contract_metadata(document(text)).get("contract_title").is_missing


def test_explicit_short_title_retained_for_review():
    field = extract_contract_metadata(document("合同名称：采购合同")).get("contract_title")
    assert field.field_value == "采购合同"
    assert field.confidence == 0.3


def test_default_missing_title_no_fabricated_anchor():
    field = extract_contract_metadata(document("甲方：示例公司")).get("contract_title")
    assert field.is_missing and field.anchor is None


def test_task_name_fallback_is_auto_without_document_anchor():
    result = extract_contract_metadata(document("甲方：示例公司"), task_name="  采购审批任务  ")
    field = result.get("contract_title")
    assert field.field_value == "采购审批任务"
    assert field.source == "auto"
    assert field.confidence == 0.3
    assert field.anchor is None
    assert field.warnings


def test_real_title_wins_over_task_name():
    field = extract_contract_metadata(document("设备采购合同"), task_name="采购审批任务").get(
        "contract_title"
    )
    assert field.field_value == "设备采购合同"
    assert field.anchor is not None


def test_overlong_fields_do_not_block_extraction():
    result = extract_contract_metadata(
        document("合同名称：" + "设备" * 300 + "采购合同", "合同编号：" + "A" * 600),
        task_name="T" * 513,
    )
    assert result.get("contract_title").is_missing
    assert result.get("contract_no").is_missing
    for key in ("contract_title", "contract_no"):
        assert len(result.get(key).field_value_raw) <= 1024
        assert json.loads(result.get(key).field_value_raw)


def test_candidate_json_stays_bounded():
    field = extract_contract_metadata(document(*(f"合同编号：HT-{i:04d}" for i in range(30)))).get(
        "contract_no"
    )
    assert len(field.field_value_raw) <= 1024
    assert json.loads(field.field_value_raw)


def test_real_docx_full_metadata_merge(tmp_path):
    word = Document()
    for text in [
        "设备采购合同",
        "合同编号：HT-2026/001",
        "甲方：示例采购有限公司",
        "合同金额：￥1000元",
        "合同期限：2026年1月1日至2026年1月31日",
        "本合同自双方盖章之日起生效。",
        "签署日期：2026年1月1日",
    ]:
        word.add_paragraph(text)
    path = tmp_path / "contract.docx"
    word.save(path)
    parsed = parse_docx(path, file_id=12)
    result = extract_metadata(parsed, today=date(2026, 10, 15))
    for key, value in [
        ("contract_title", "设备采购合同"),
        ("contract_no", "HT-2026/001"),
        ("party_a_name", "示例采购有限公司"),
        ("amount", "1000.00"),
        ("contract_period_days", "30"),
        ("sign_date", "2026-01-01"),
    ]:
        assert result.get(key).field_value == value
    assert len(result.items) == 15
    result.validate_against_document(parsed)
