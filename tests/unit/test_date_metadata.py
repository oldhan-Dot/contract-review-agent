"""日期字段的上下文、非法日期、期限互校与真实 DOCX 证据验证。"""

import json
from datetime import date

import pytest
from docx import Document

from app.parsing.docx_channel import parse_docx
from app.parsing.ir import ParsedDocument
from app.structuring.date_metadata import extract_date_metadata
from app.structuring.metadata import extract_metadata

TODAY = date(2026, 10, 15)


def document(*texts):
    paragraphs, cursor = [], 0
    for index, text in enumerate(texts):
        paragraphs.append(
            dict(
                pid=f"file-12:p{index}",
                index=index,
                text=text,
                char_start=cursor,
                char_end=cursor + len(text),
            )
        )
        cursor += len(text) + 1
    full_text = "\n".join(texts)
    return ParsedDocument(
        doc_id="file-12",
        file_id=12,
        source_format="DOCX",
        channel="docx",
        paragraphs=paragraphs,
        full_text=full_text,
        quality=dict(precision_cap="EXACT", char_count=len(full_text), paragraph_count=len(texts)),
    )


def extract(*texts, today=TODAY):
    return extract_date_metadata(document(*texts), today=today)


@pytest.mark.parametrize(
    "start,end",
    [
        ("2026年1月1日", "2026年1月31日"),
        ("2026-01-01", "2026-01-31"),
        ("2026/1/1", "2026/1/31"),
        ("2026.1.1", "2026.1.31"),
        ("２０２６年１月１日", "２０２６年１月３１日"),
        ("2026 年 1 月 1 日", "2026 年 1 月 31 日"),
    ],
)
def test_date_range_formats_and_exact_anchors(start, end):
    parsed = document(f"合同期限：自{start}起至{end}止")
    result = extract_date_metadata(parsed, today=TODAY)
    assert result.get("perform_start_date").field_value == "2026-01-01"
    assert result.get("perform_end_date").field_value == "2026-01-31"
    assert result.get("contract_period_days").field_value == "30"
    assert result.get("perform_start_date").anchor.matched_text == start
    result.validate_against_document(parsed)


def test_separate_labels_and_multiline_value():
    result = extract("履行开始日期：", "2026-01-01", "履行结束日期：2026-01-31")
    assert result.get("perform_start_date").field_value == "2026-01-01"
    assert result.get("contract_period_days").field_value == "30"


def test_open_date_does_not_fabricate_other_endpoint():
    result = extract("履行期限自2026年1月1日起，期限30日")
    assert result.get("perform_start_date").field_value == "2026-01-01"
    assert result.get("perform_end_date").is_missing
    assert result.get("contract_period_days").field_value == "30"


def test_end_only():
    result = extract("履行期限至2026年12月31日止")
    assert result.get("perform_end_date").field_value == "2026-12-31"
    assert result.get("perform_start_date").is_missing


@pytest.mark.parametrize("raw", ["2026年2月29日", "2026-13-01", "2026/4/31", "0000年1月1日"])
def test_invalid_calendar_dates_retained_for_review(raw):
    field = extract(f"履行开始日期：{raw}").get("perform_start_date")
    assert field.field_value == raw
    assert field.confidence == 0.3
    assert field.warnings
    assert field.anchor.matched_text == raw


def test_leap_date_is_valid():
    field = extract("履行开始日期：2024年2月29日").get("perform_start_date")
    assert field.field_value == "2024-02-29"
    assert field.confidence == 0.9


def test_reversed_dates_warn_without_sorting_or_negative_period():
    result = extract("合同期限：2026-02-01至2026-01-01")
    assert result.get("perform_start_date").field_value == "2026-02-01"
    assert result.get("perform_end_date").field_value == "2026-01-01"
    assert result.get("perform_start_date").confidence == 0.3
    assert result.get("contract_period_days").is_missing


@pytest.mark.parametrize(
    "text",
    [
        "付款日期：2026年1月1日，交付日期：2026年2月1日",
        "本合同付款自2026年1月1日至2026年2月1日",
        "收到发票后30日付款，质保期365日，提前15日通知",
        "履行开始日期：2026-01/01",
        "履行开始日期：20260年1月1日",
        "签署日期：____年__月__日",
    ],
)
def test_unrelated_or_incomplete_dates_not_guessed(text):
    result = extract(text)
    for key in ("perform_start_date", "perform_end_date", "sign_date", "contract_period_days"):
        assert result.get(key).is_missing


@pytest.mark.parametrize(
    "text,expected",
    [
        ("合同期限为30个自然日", "30"),
        ("合同期限：90天", "90"),
        ("本合同自生效之日起60日内履行完毕", "60"),
        ("合同期限：2024年1月1日至2025年1月1日，为1年", "366"),
        ("履行期限自2024年2月1日起，共1个月", "29"),
    ],
)
def test_explicit_period_and_calendar_conversion(text, expected):
    field = extract(text).get("contract_period_days")
    assert field.field_value == expected
    assert field.warnings == []


@pytest.mark.parametrize(
    "text",
    [
        "合同期限为1个月",
        "合同期限为1年",
        "合同期限为30个工作日",
        "履行期限自2026年1月31日起，共1个月",
    ],
)
def test_month_year_and_workdays_not_approximated(text):
    field = extract(text).get("contract_period_days")
    assert field.is_missing
    assert field.warnings


def test_period_conflict_uses_date_difference_and_preserves_raw():
    field = extract("合同期限：2026-01-01至2026-01-31，共90日").get("contract_period_days")
    assert field.field_value == "30"
    assert field.confidence == 0.3
    assert "不一致" in "".join(field.warnings)
    assert any(item["raw"] == "90日" for item in json.loads(field.field_value_raw))


def test_zero_explicit_duration_is_retained_as_invalid():
    field = extract("合同期限为0天").get("contract_period_days")
    assert field.field_value == "0"
    assert field.confidence == 0.3


def test_multiple_candidates_preserved_with_warning():
    result = extract("履行开始日期：2026-01-01", "履行开始日期：2026-02-01")
    field = result.get("perform_start_date")
    assert field.field_value == "2026-01-01"
    assert len(json.loads(field.field_value_raw)) == 2
    assert field.confidence == 0.3


def test_signature_last_contextual_date_not_later_delivery_date():
    result = extract("签署日期：2026年1月1日", "签署日期：2026年1月2日", "交付日期：2026年10月1日")
    field = result.get("sign_date")
    assert field.field_value == "2026-01-02"
    assert len(json.loads(field.field_value_raw)) == 2
    assert field.warnings


def test_signature_block_next_date_line():
    field = extract("甲方（盖章）：示例公司", "日期：2026年10月1日").get("sign_date")
    assert field.field_value == "2026-10-01"
    assert field.anchor.paragraph_index == 1


@pytest.mark.parametrize(
    "raw,warning", [("2021-10-14", True), ("2021-10-15", False), ("2026-10-01", False)]
)
def test_signature_five_year_boundary(raw, warning):
    field = extract(f"签署日期：{raw}").get("sign_date")
    assert bool(field.warnings) is warning
    assert field.field_value == raw


def test_signature_five_year_leap_day_boundary():
    field = extract("签署日期：2019年2月28日", today=date(2024, 2, 29)).get("sign_date")
    assert field.warnings == []


def test_effective_condition_keeps_exact_paragraph():
    text = "本合同自双方签字盖章之日起生效。"
    field = extract(text).get("effective_condition")
    assert field.field_value == text
    assert field.anchor.matched_text == text


def test_long_effective_paragraph_keeps_keyword_and_real_position():
    parsed = document("说明" * 180 + "本合同自双方盖章后生效。" + "后续" * 180)
    field = extract_date_metadata(parsed, today=TODAY).get("effective_condition")
    assert len(field.field_value) == 200
    assert "生效" in field.field_value
    assert field.confidence == 0.3
    field.anchor.validate_against_document(parsed)


def test_large_candidate_set_is_bounded_valid_json():
    field = extract(*(f"履行开始日期：2026-01-{day:02d}" for day in range(1, 29))).get(
        "perform_start_date"
    )
    assert len(field.field_value_raw) <= 1024
    assert json.loads(field.field_value_raw)


def test_real_docx_combines_all_completed_extractors(tmp_path):
    word = Document()
    for text in [
        "甲方：示例采购有限公司",
        "合同金额：人民币1000元",
        "合同期限：2026年1月1日至2026年1月31日",
        "本合同自双方签字盖章之日起生效。",
        "签署日期：2026年1月1日",
    ]:
        word.add_paragraph(text)
    path = tmp_path / "date.docx"
    word.save(path)
    parsed = parse_docx(path, file_id=12)
    result = extract_metadata(parsed, today=TODAY)
    for key, expected in [
        ("party_a_name", "示例采购有限公司"),
        ("amount", "1000.00"),
        ("perform_start_date", "2026-01-01"),
        ("contract_period_days", "30"),
        ("sign_date", "2026-01-01"),
    ]:
        assert result.get(key).field_value == expected
    assert result.get("contract_no").is_missing
    result.validate_against_document(parsed)


@pytest.mark.parametrize(
    "text", ["报价有效期：2026年1月1日至2026年1月31日", "报价有效期为30天", "质保有效期为365日"]
)
def test_component_validity_not_contract_period(text):
    result = extract(text)
    assert result.get("perform_start_date").is_missing
    assert result.get("contract_period_days").is_missing


@pytest.mark.parametrize("raw", ["-30", "－３０"])
def test_negative_period_not_silently_changed_to_positive(raw):
    field = extract(f"合同期限为{raw}天").get("contract_period_days")
    assert field.field_value == "-30"
    assert field.confidence == 0.3
    assert field.anchor.matched_text == f"{raw}天"


@pytest.mark.parametrize("raw", ["1个月", "30个工作日"])
def test_unconverted_duration_still_preserves_raw(raw):
    field = extract(f"合同期限为{raw}").get("contract_period_days")
    assert field.is_missing
    assert json.loads(field.field_value_raw)[0]["raw"] == raw
