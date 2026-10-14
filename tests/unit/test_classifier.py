"""使用独立条款片段验证加权归类、证据、继承和失败兜底。"""

import json

import pytest
from docx import Document

from app.parsing.docx_channel import parse_docx
from app.structuring.classifier import classify_clause, classify_clauses
from app.structuring.models import Clause, ClauseType
from app.structuring.splitter import split_clauses


def clause(text, *, title=None, index=1, parent_index=None):
    return Clause(
        index=index,
        title=title,
        text=text,
        level=1,
        paragraph_index_start=0,
        paragraph_index_end=0,
        char_start=0,
        char_end=len(text),
        split_strategy="NUMBER",
        parent_index=parent_index,
    )


@pytest.mark.parametrize(
    "text,expected",
    [
        ("甲方：示例科技有限公司", "PARTY_INFO"),
        ("货物名称为工业设备。", "SUBJECT_MATTER"),
        ("合同价款应按约定支付。", "PAYMENT"),
        ("货物应达到验收标准。", "DELIVERY_ACCEPTANCE"),
        ("逾期应承担违约责任并支付违约金。", "LIABILITY"),
        ("不得披露商业秘密。", "CONFIDENTIALITY"),
        ("知识产权及著作权归采购方。", "IP"),
        ("应遵守个人信息和数据保护要求。", "DATA_SECURITY"),
        ("争议提交人民法院诉讼解决。", "DISPUTE"),
        ("合同有效期自双方签字后开始。", "TERM_EFFECTIVE"),
        ("提前终止应交接已完成的工作。", "TERMINATION"),
        ("因自然灾害或战争无法履行时适用不可抗力约定。", "FORCE_MAJEURE"),
    ],
)
def test_all_twelve_types_have_expected_body_classification(text, expected):
    result = classify_clause(clause(text))
    assert result.clause_type == expected
    assert result.type_confidence > 0
    evidence = json.loads(result.classify_rule_hit)
    assert evidence["method"] == "keywords"
    assert expected in evidence["hits"]


def test_title_overrides_body_and_input_is_not_changed():
    original = clause("损失应由责任方赔偿。", title="付款约定")
    result = classify_clause(original)
    assert result.clause_type == ClauseType.PAYMENT
    assert original.clause_type == ClauseType.OTHER
    assert result.char_start == original.char_start and result.char_end == original.char_end
    assert json.loads(result.classify_rule_hit)["method"] == "title"


def test_subject_title_alias_is_recognized():
    assert classify_clause(clause("双方信息。", title="主体信息")).clause_type == "PARTY_INFO"


@pytest.mark.parametrize("title", ["定义与解释", "通知与送达"])
def test_general_sections_are_explicit_other(title):
    result = classify_clause(clause("本节包含付款和赔偿等术语的解释。", title=title))
    assert result.clause_type == ClauseType.OTHER
    assert json.loads(result.classify_rule_hit)["needs_fallback"] is False


def test_mentioning_parties_does_not_override_payment_topic():
    result = classify_clause(clause("甲方应向乙方支付价款。"))
    assert result.clause_type == ClauseType.PAYMENT
    assert "PARTY_INFO" not in json.loads(result.classify_rule_hit)["scores"]


def test_repeated_keyword_is_scored_once():
    once = json.loads(classify_clause(clause("保密。")).classify_rule_hit)["scores"]
    repeated = json.loads(classify_clause(clause("保密。保密。保密。")).classify_rule_hit)["scores"]
    assert repeated == once


def test_close_second_type_is_recorded_without_duplicate_clause():
    result = classify_clause(clause("付款需按约到货。"))
    assert result.clause_type == "PAYMENT"
    assert result.clause_type_secondary == "DELIVERY_ACCEPTANCE"


@pytest.mark.parametrize(
    "text,reason",
    [
        ("付款和保密。", "tied_scores"),
        ("通过汇款办理。", "low_score"),
        ("按约定办理。", "no_keywords"),
    ],
)
def test_unresolved_rules_use_other_and_keep_reason(text, reason):
    result = classify_clause(clause(text))
    assert result.clause_type == ClauseType.OTHER
    assert result.type_confidence == 0
    assert json.loads(result.classify_rule_hit)["reason"] == reason
    assert json.loads(result.classify_rule_hit)["needs_fallback"] is True


def test_child_inherits_parent_when_own_text_has_no_decisive_topic():
    parent = clause("按约定办理。", title="付款约定")
    child = clause("应于三日内办理。", index=2, parent_index=1)
    result = classify_clauses([parent, child])
    assert result[1].clause_type == ClauseType.PAYMENT
    assert json.loads(result[1].classify_rule_hit)["method"] == "parent"


def test_general_definition_children_remain_other():
    parent = clause("以下术语的定义。", title="定义与解释")
    child = clause("付款的含义如下。", title="付款", index=2, parent_index=1)
    result = classify_clauses([parent, child])
    assert result[1].clause_type == "OTHER"
    assert json.loads(result[1].classify_rule_hit)["method"] == "general_parent"


def test_stronger_child_body_can_override_parent_type():
    parent = clause("付款。")
    child = clause("数据安全和个人信息保护。", index=2, parent_index=1)
    result = classify_clauses([parent, child])
    assert result[1].clause_type == ClauseType.DATA_SECURITY


def test_missing_parent_is_rejected():
    with pytest.raises(ValueError, match="父条款"):
        classify_clauses([clause("付款。", index=2, parent_index=1)])


def test_fallback_only_runs_when_needed():
    calls = []

    def fallback(current):
        calls.append(current.index)
        return {"clause_type": "CONFIDENTIALITY", "confidence": 0.8}

    classified = classify_clauses(
        [clause("付款。"), clause("妥善保管所获资料。", index=2)], fallback=fallback
    )
    assert calls == [2]
    assert classified[1].clause_type == "CONFIDENTIALITY"
    assert json.loads(classified[1].classify_rule_hit)["method"] == "fallback"


@pytest.mark.parametrize(
    "answer",
    [
        {"clause_type": "MADE_UP", "confidence": 0.8},
        {"clause_type": "PAYMENT", "confidence": 2},
        {"clause_type": "PAYMENT", "confidence": float("nan")},
        {"clause_type": "PAYMENT", "confidence": True},
        {"clause_type": "PAYMENT"},
    ],
)
def test_invalid_fallback_output_degrades_to_other(answer):
    result = classify_clause(clause("按约定办理。"), fallback=lambda current: answer)
    assert result.clause_type == "OTHER"
    assert json.loads(result.classify_rule_hit)["needs_fallback"] is True


def test_fallback_failure_does_not_expose_error_text():
    def failing(current):
        raise RuntimeError("模拟敏感错误内容")

    result = classify_clause(clause("按约定办理。"), fallback=failing)
    assert result.clause_type == "OTHER"
    assert "模拟敏感错误内容" not in result.classify_rule_hit


def test_fallback_cannot_mutate_callers_clause():
    original = clause("按约定办理。")

    def fallback(current):
        current.text = "修改后的内容"
        return {"clause_type": "PAYMENT", "confidence": 0.8}

    classified = classify_clause(original, fallback=fallback)
    assert original.text == "按约定办理。"
    assert classified.text == original.text


def test_real_docx_can_be_parsed_split_and_classified_with_original_ranges(tmp_path):
    word = Document()
    word.add_paragraph("设备采购合同")
    word.add_heading("第一条 付款", level=1)
    word.add_paragraph("甲方应向乙方支付合同价款，并在付款时取得合法有效的发票。" * 2)
    word.add_heading("第二条 保密", level=1)
    word.add_paragraph("双方不得披露履约中获悉的商业秘密。")
    path = tmp_path / "contract.docx"
    word.save(path)
    document = parse_docx(path, file_id=1)
    original = split_clauses(document)
    result = classify_clauses(original)
    assert [c.clause_type for c in result] == ["OTHER", "PAYMENT", "CONFIDENTIALITY"]
    for before, after in zip(original, result):
        assert before.clause_type == "OTHER"
        assert before.source_text(document) == after.source_text(document)
        after.validate_against_document(document)
